"""Slide visuals: turn a deck PDF into page images, or pull the pictures out of the .pptx, so Claude can see them.
Images are shrunk and saved next to the session; the pre-read sends them with the slide text."""
from __future__ import annotations

import base64
import io
from pathlib import Path

MAX_SLIDES = 60        # cost guard: images beyond this are skipped
PAGE_WIDTH = 1100      # px, enough to read chart labels
PER_SLIDE = 3          # embedded pictures kept per slide when there is no PDF


def _jpeg(img, width: int) -> bytes:
    from PIL import Image
    img = img.convert("RGB")
    if img.width > width:
        img = img.resize((width, max(1, round(img.height * width / img.width))), Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=72, optimize=True)
    return buf.getvalue()


def render_pdf(data: bytes) -> list[bytes]:
    """One JPEG per PDF page (a deck exported from PowerPoint: page N is slide N)."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(data)
    out = []
    for i in range(min(len(pdf), MAX_SLIDES)):
        page = pdf[i]
        w = page.get_width() or 960
        out.append(_jpeg(page.render(scale=PAGE_WIDTH / w).to_pil(), PAGE_WIDTH))
    return out


def _pictures(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from _pictures(sh.shapes)
        else:
            try:
                blob = sh.image.blob          # picture, or a placeholder holding one
            except Exception:
                continue
            yield blob


def pptx_pictures(data: bytes) -> dict[int, list[bytes]]:
    """Meaningful embedded pictures per slide (logos and tiny icons are dropped; EMF/WMF cannot be read)."""
    from PIL import Image
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    res: dict[int, list[bytes]] = {}
    for n, slide in enumerate(prs.slides, 1):
        if n > MAX_SLIDES:
            break
        found = []
        for blob in _pictures(slide.shapes):
            try:
                im = Image.open(io.BytesIO(blob)); im.load()
            except Exception:
                continue
            if min(im.size) < 120 or im.width * im.height < 40000:
                continue
            found.append((im.width * im.height, im))
        found.sort(key=lambda x: -x[0])
        if found:
            res[n] = [_jpeg(im, 1000) for _, im in found[:PER_SLIDE]]
    return res


def save(folder: Path, images: dict[int, list[bytes]]) -> int:
    folder.mkdir(parents=True, exist_ok=True)
    for f in folder.glob("*.jpg"):
        f.unlink()
    n = 0
    for slide, lst in images.items():
        for k, b in enumerate(lst, 1):
            (folder / f"s{int(slide):03d}_{k}.jpg").write_bytes(b); n += 1
    return n


def load(folder: Path) -> list[dict]:
    """[{n, k, b64}] in slide order, for the model."""
    out = []
    for f in sorted(folder.glob("s*_*.jpg")) if folder.exists() else []:
        try:
            n, k = f.stem[1:].split("_")
            out.append({"n": int(n), "k": int(k), "b64": base64.b64encode(f.read_bytes()).decode()})
        except Exception:
            continue
    return out
