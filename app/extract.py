"""Pull plain text out of the uploaded deck (.pptx) and RFP (.pdf/.docx/.txt/.md)."""
from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu


def _shape_text(shape) -> list[str]:
    out: list[str] = []
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            out.extend(_shape_text(s))
        return out
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            out.append(t)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    return out


def _reading_order(shapes):
    """Top-to-bottom, left-to-right, so text reads the way the audience sees it."""
    def key(s):
        top = s.top if s.top is not None else Emu(0)
        left = s.left if s.left is not None else Emu(0)
        return (round(int(top) / 300000), int(left))
    return sorted(shapes, key=key)


def extract_deck(data: bytes) -> list[dict]:
    prs = Presentation(io.BytesIO(data))
    slides: list[dict] = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        texts: list[str] = []
        pictures = 0
        for shape in _reading_order(slide.shapes):
            if shape.shape_type == 13:
                pictures += 1
            texts.extend(_shape_text(shape))
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if not title and texts:
            title = texts[0].splitlines()[0][:90]
        title = " – ".join(x.strip() for x in title.replace("\x0b", "\n").splitlines() if x.strip())
        texts = [t.replace("\x0b", "\n") for t in texts]
        slides.append(
            {"n": i, "title": title, "text": "\n".join(texts), "notes": notes, "pictures": pictures}
        )
    return slides


def extract_document(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        import pdfplumber

        parts = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for n, page in enumerate(pdf.pages, 1):
                t = (page.extract_text() or "").strip()
                if t:
                    parts.append(f"[page {n}]\n{t}")
        return "\n\n".join(parts)
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for tbl in d.tables:
            for row in tbl.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines)
    if ext in {".txt", ".md", ".text", ".vtt", ".srt"}:
        return data.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported RFP file type: {ext or 'unknown'} (use PDF, DOCX, TXT or MD)")


def deck_as_prompt_text(slides: list[dict]) -> str:
    blocks = []
    for s in slides:
        b = [f"=== SLIDE {s['n']}: {s['title']}", "ON SLIDE:", s["text"] or "(no text; image-only slide)"]
        if s["notes"]:
            b += ["SPEAKER NOTES (presenter prompts, not shown to the audience):", s["notes"]]
        blocks.append("\n".join(b))
    return "\n\n".join(blocks)
