"""Word (.docx) reports: the pre-read (slide map, requirements, questions, gaps) and rehearsal results."""
from __future__ import annotations

import io
import time

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x8A, 0x1C, 0x2B)
PERSONA_LABEL = {"finance": "Finance", "operations": "Operations", "skeptic": "Skeptic",
                 "insights": "Insights", "channel_expert": "Channel expert"}


# ---------- helpers ----------
def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _set_text(cell, text, bold=False, size=9, color=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text if text is not None else ""))
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def _table(doc, headers, rows, widths_cm, fills=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        _set_text(c, h, bold=True, size=9, color=RGBColor(255, 255, 255))
        _shade(c, "8A1C2B")
        c.width = Cm(widths_cm[i])
    _repeat_header(t.rows[0])
    for r_i, r in enumerate(rows):
        cells = t.add_row().cells
        _no_split(t.rows[-1])
        for i, v in enumerate(r):
            _set_text(cells[i], v, size=9)
            cells[i].width = Cm(widths_cm[i])
            if fills and fills[r_i] and fills[r_i].get(i):
                _shade(cells[i], fills[r_i][i])
    for i, w in enumerate(widths_cm):  # also set the grid so LibreOffice/Word honour the widths
        t.columns[i].width = Cm(w)
    return t


def _h(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = ACCENT
    return h


def _para(doc, text, size=10, bold=False, italic=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return p


def _new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.6))
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    return doc


def _bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


FILL = {"strong": "DDEFE3", "covered": "DDEFE3", "ok": "EEECE6", "thin": "FBEBC8", "none": "F6D5D2",
        "unanswered": "F6D5D2", "ignored": "EEECE6"}


def req_status(r: dict, slides: list) -> str:
    if r.get("ignored"):
        return "ignored"
    cov = [s for s in slides if r["id"] in s.get("covers", [])]
    if not cov:
        return "unanswered"
    return "thin" if all(s["evidence_strength"] in ("thin", "none") for s in cov) else "covered"


# ---------- pre-read report ----------
def _preread_sections(doc, sess: dict, pr: dict) -> None:
    st = sess["settings"]
    ov = sess.get("pacing_overrides", {})
    slides = pr["slides"]
    titles = {s["n"]: s["title"] for s in sess["slides"]}
    reqs = pr["requirements"]
    status = {r["id"]: req_status(r, slides) for r in reqs}
    cnt = lambda k: sum(1 for v in status.values() if v == k)

    _h(doc, "At a glance", 1)
    _para(doc, pr.get("deck_summary", ""), 11)
    _table(doc, ["Requirements covered", "Thin evidence", "Unanswered", "Inconsistencies", "Likely questions"],
           [[cnt("covered"), cnt("thin"), cnt("unanswered"), len(pr["inconsistencies"]),
             sum(len(s["likely_questions"]) for s in slides)]], [5, 5, 5, 5, 5])
    out = (st.get("ignore_topics") or "").strip()
    if out:
        _para(doc, f"Out of scope for this review: {out}.", 9, italic=True)

    _h(doc, "1. Slide map", 1)
    _para(doc, "Purpose = the job the slide does in the argument. Pacing is the intended speed (quick / normal / weighty); "
               "entries marked * were overridden by the presenter.", 9, italic=True)
    rows, fills = [], []
    for s in slides:
        pace = ov.get(str(s["n"])) or s["pacing"]
        star = "*" if str(s["n"]) in ov else ""
        rows.append([s["n"], titles.get(s["n"], ""), f"{s['purpose']}\nKey message: {s['key_message']}",
                     pace + star, f"{s['evidence_strength']}: {s['evidence_note']}", ", ".join(s["covers"])])
        fills.append({4: FILL.get(s["evidence_strength"])})
    _table(doc, ["#", "Slide", "Purpose and key message", "Pacing", "Evidence", "Covers"], rows,
           [1.0, 4.2, 9.2, 2.0, 7.6, 2.4], fills)

    _h(doc, "2. Requirements", 1)
    rows, fills = [], []
    for r in reqs:
        cover = [str(s["n"]) for s in slides if r["id"] in s["covers"]]
        k = status[r["id"]]
        rows.append([r["id"], r["text"], r["priority"], r["category"],
                     "out of scope" if k == "ignored" else k, ", ".join(cover) or "–"])
        fills.append({4: FILL.get(k)})
    _table(doc, ["ID", "Requirement", "Priority", "Category", "Status", "Slides"], rows,
           [1.2, 13.0, 2.0, 3.6, 2.6, 4.0], fills)

    _h(doc, "3. Likely questions", 1)
    by = {}
    for s in slides:
        for q in s["likely_questions"]:
            by.setdefault(q["persona_hint"], []).append((s["n"], q["question"]))
    for k, label in PERSONA_LABEL.items():
        if k in by:
            _h(doc, f"{label} ({len(by[k])})", 2)
            _table(doc, ["Slide", "Question"], [[n, q] for n, q in by[k]], [1.8, 24.0])

    _h(doc, "4. Gaps and inconsistencies", 1)
    _h(doc, "Unanswered or only touched on", 2)
    rmap = {r["id"]: r for r in reqs}
    if pr["unanswered_requirements"]:
        _table(doc, ["ID", "Requirement", "Why it's a gap"],
               [[u["req_id"], rmap[u["req_id"]]["text"], u["note"]] for u in pr["unanswered_requirements"]],
               [1.4, 11.0, 13.4])
    else:
        _para(doc, "None found.")
    _h(doc, "Inconsistencies", 2)
    if pr["inconsistencies"]:
        _table(doc, ["Slides", "What doesn't add up"],
               [[", ".join(str(n) for n in i["slides"]), i["description"]] for i in pr["inconsistencies"]], [3.0, 22.8])
    else:
        _para(doc, "None found.")


def _title_block(doc, sess: dict, subtitle: str) -> None:
    t = doc.add_paragraph()
    r = t.add_run(sess["settings"].get("title") or "Pitch review")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = ACCENT
    _para(doc, subtitle, 11, color=RGBColor(0x6B, 0x6A, 0x66))
    _para(doc, f"Prepared {time.strftime('%d %B %Y')} · AI-generated review of the deck against the brief. "
               "Treat it as a prompt for discussion and check any figures against the source documents.", 9, italic=True)


def build_preread_docx(sess: dict, pr: dict) -> bytes:
    doc = _new_doc()
    _title_block(doc, sess, "Pre-pitch review: slide map, requirements, likely questions and gaps")
    _preread_sections(doc, sess, pr)
    return _bytes(doc)


# ---------- rehearsal report ----------
def _score_sections(doc, reh: dict) -> None:
    sc = reh["score"]
    _h(doc, "Rehearsal results", 1)
    _para(doc, f"Rehearsal {reh.get('started', '')} · {(str(sc['meta'].get('minutes')) + ' min pitch') if sc['meta'].get('minutes') else 'timing not assessed'}", 9, italic=True)
    _table(doc, ["Brief coverage", "Gut feel", "Question handling", "Delivery", "Total (coverage + gut)"],
           [[sc["brief_coverage"]["score"], sc["gut_feel"]["score"], sc["question_handling"]["score"],
             "n/a" if sc["delivery_score"] is None else sc["delivery_score"], sc["total"]]], [5, 5, 5, 5, 5])
    w = sc["win_likelihood"]
    _para(doc, f"Win likelihood: panel's independent read {w['panel_read_pct']}% vs calculated {w['calculated_pct']}% "
               + ("(0.41 coverage + 0.35 gut + 0.24 questions; no delivery score without timestamps)." if sc["delivery_score"] is None
                else "(0.35 coverage + 0.30 gut + 0.20 questions + 0.15 delivery)."), 10, bold=True)
    _para(doc, w.get("panel_rationale", ""), 10)

    _h(doc, "Top actions", 2)
    for a in sc["top_actions"]:
        doc.add_paragraph(a, style="List Bullet")

    _h(doc, "Brief coverage", 2)
    _para(doc, sc["brief_coverage"].get("summary", ""))
    _table(doc, ["Req", "Met?", "Evidence from the pitch", "Note"],
           [[i["req_id"], i["met"], i["evidence"], i["note"]] for i in sc["brief_coverage"]["items"]],
           [1.4, 2.0, 11.0, 11.4],
           [{1: {"yes": "DDEFE3", "partial": "FBEBC8", "no": "F6D5D2"}.get(i["met"])} for i in sc["brief_coverage"]["items"]])

    _h(doc, "Gut feel", 2)
    _para(doc, sc["gut_feel"].get("rationale", ""))

    _h(doc, "Question handling", 2)
    _table(doc, ["Asked by", "Question", "How it was answered", "Rating", "Note"],
           [[i["asker"], i["question"], i["answer_summary"], i["rating"], i["note"]] for i in sc["question_handling"]["items"]],
           [3.2, 7.0, 7.0, 1.6, 7.0])

    _h(doc, "Slide-by-slide delivery", 2)
    _table(doc, ["#", "Slide", "Intended pace", "Time (s)", "Words", "Verdict", "Rating", "Note"],
           [[s["n"], s["title"], s["pacing"], "" if s["seconds"] is None else s["seconds"], s["words"], s["verdict"], s["rating"], s["note"]]
            for s in sc["slides"]], [1.0, 4.4, 2.0, 1.6, 1.4, 2.0, 1.4, 12.0],
           [{5: {"clear": "DDEFE3", "ok": "EEECE6", "rushed": "FBEBC8", "waffled": "FBEBC8", "skipped": "F6D5D2"}.get(s["verdict"])}
            for s in sc["slides"]])

    if sc["presenters"]:
        _h(doc, "Per presenter", 2)
        for p in sc["presenters"]:
            _para(doc, f"{p['name']} (slides {p['slides'] or '–'})", 11, bold=True)
            for x in p["strengths"]:
                doc.add_paragraph("Strength: " + x, style="List Bullet")
            for x in p["improve"]:
                doc.add_paragraph("Improve: " + x, style="List Bullet")


def build_rehearsal_docx(sess: dict, pr: dict, reh: dict, include_preread: bool = True) -> bytes:
    doc = _new_doc()
    _title_block(doc, sess, "Rehearsal report" + (" and pre-pitch review" if include_preread else ""))
    if reh.get("past"):
        pa = reh["past"]
        doc.add_paragraph("Assessed from a transcript of a real pitch" + (f" (actual outcome: {pa['outcome']})" if pa.get("outcome") else "")
                          + (". The recording had no timestamps, so timing and delivery are not scored." if pa.get("timing_estimated") else "."))
    _score_sections(doc, reh)
    if include_preread and pr:
        doc.add_page_break()
        _preread_sections(doc, sess, pr)
    return _bytes(doc)


# ---------- post-pitch review (real pitch, from a transcript) ----------
VERDICT_TEXT = {"covered": "Covered", "partly": "Partly covered", "deck_only": "On slides only", "conversation_only": "Said, not on slides", "gap": "Gap"}
VERDICT_FILL = {"covered": "DDEFE3", "partly": "FBEBC8", "deck_only": "FBEBC8", "conversation_only": "FBEBC8", "gap": "F6D5D2"}
ANS_TEXT = {"yes": "Answered", "partly": "Partly", "no": "Not answered", "deferred": "Deferred"}
ANS_FILL = {"yes": "DDEFE3", "partly": "FBEBC8", "deferred": "FBEBC8", "no": "F6D5D2"}


def _review_sections(doc, reh: dict) -> None:
    rv = reh["review"]
    c = rv["counts"]
    bp = c["by_priority"]
    pa = reh.get("past") or {}
    _para(doc, "Review of the pitch as delivered" + (f" · outcome: {pa['outcome']}" if pa.get("outcome") else ""), 9, italic=True)
    _h(doc, "Summary", 1)
    n = c["requirements"]
    _table(doc, ["Addressed in the conversation", "Supported by the deck", "Covered by both", "Must-haves covered", "Should-haves covered", "Pitch-level gaps", "Client questions answered"],
           [[f"{c['conv_addressed']} / {n} ({c['conv_delivered']} fully)", f"{c['deck_supported']} / {n} ({c['deck_strong']} strongly)", f"{c['covered']} / {n}",
             f"{bp['must']['covered']} / {bp['must']['total']}", f"{bp['should']['covered']} / {bp['should']['total']}",
             f"{c['pitch_gaps']} (+{c['contract_gaps']} contract detail)", f"{c['q_yes']} / {c['questions']}"]],
           [3.8, 3.8, 3.2, 3.2, 3.2, 3.8, 3.4])
    _para(doc, "Covered = the deck supports it at least partly AND the conversation delivered it. Gaps on contract-level detail are not a failing of the pitch itself.", 8, italic=True)
    _para(doc, rv["summary"])
    if rv["biggest_gaps"]:
        _h(doc, "Biggest gaps", 2)
        for g in rv["biggest_gaps"]:
            doc.add_paragraph(g, style="List Bullet")
    if rv["next_steps"]:
        _h(doc, "Next steps", 2)
        for g in rv["next_steps"]:
            doc.add_paragraph(g, style="List Bullet")

    _h(doc, "By theme", 1)
    _table(doc, ["Theme", "Result", "Requirements", "Addressed in conversation", "Supported by deck", "Gaps", "How the pitch did"],
           [[t["name"], t["status"].title(), t["counts"]["total"], t["counts"]["conv_addressed"], t["counts"]["deck_supported"], t["counts"]["gap"], t["summary"]]
            for t in rv["themes"]], [4.0, 1.8, 2.2, 2.6, 2.4, 1.4, 12.0],
           [{1: {"strong": "DDEFE3", "mixed": "FBEBC8", "weak": "F6D5D2"}.get(t["status"])} for t in rv["themes"]])

    _h(doc, "Requirements: deck and conversation", 1)
    by_id = {r["req_id"]: r for r in rv["requirements"]}
    rows, fills = [], []
    for t in rv["themes"]:
        for rid in t["req_ids"]:
            r = by_id[rid]
            rows.append([r["req_id"],
                         f"{r['text']} ({r['priority']}{', contract / onboarding detail' if r['relevance'] == 'contract' else ''}) · theme: {t['name']}",
                         f"{r['deck']['strength'].title()}" + (f" · slide{'s' if len(r['deck']['slides']) > 1 else ''} {', '.join(map(str, r['deck']['slides']))}" if r["deck"]["slides"] else "")
                         + (f"\n{r['deck']['note']}" if r["deck"]["note"] else ""),
                         {"delivered": "Delivered", "partial": "Partly", "not_discussed": "Not discussed"}[r["conversation"]["status"]]
                         + (f"\n{r['conversation']['evidence']}" if r["conversation"]["evidence"] else "")
                         + (f" ({r['conversation']['speaker']})" if r["conversation"]["speaker"] and r["conversation"]["evidence"] else "")
                         + (f"\n{r['conversation']['note']}" if r["conversation"]["note"] else ""),
                         VERDICT_TEXT[r["verdict"]], r["recommendation"]])
            fills.append({4: VERDICT_FILL.get(r["verdict"])})
    _table(doc, ["Req", "What the client asked for", "How the deck supports it", "How the conversation delivered it", "Verdict", "Recommendation"],
           rows, [1.2, 5.0, 5.4, 7.4, 2.4, 5.0], fills)

    _h(doc, "Client questions and remarks", 1)
    if rv["client_questions"]:
        _table(doc, ["#", "Asked by", "What the client said", "Relates to", "Answered?", "How it was answered", "Note / follow-up"],
               [[i + 1, q["asker"], f"{q['kind'].title()}: {q['question']}",
                 ((f"Slide {q['slide']}" if q["slide"] else "") + (f" {q['req_id']}" if q["req_id"] else "")).strip(),
                 ANS_TEXT[q["answered"]] + (f" ({q['quality']})" if q["quality"] != "n/a" else ""), q["answer_summary"],
                 " ".join(x for x in (q["note"], ("Follow-up owed: " + q["follow_up"]) if q["follow_up"] else "") if x)]
                for i, q in enumerate(rv["client_questions"])],
               [0.8, 2.4, 6.6, 2.2, 2.6, 6.0, 5.8],
               [{4: ANS_FILL.get(q["answered"])} for q in rv["client_questions"]])
    else:
        _para(doc, "No client questions or remarks were found in the transcript.")
    if rv["commitments"]:
        _h(doc, "Commitments and follow-ups owed", 2)
        _table(doc, ["What", "Who", "By when"], [[x["what"], x["who"], x["due"]] for x in rv["commitments"]], [14.0, 6.0, 6.0])


def build_review_docx(sess: dict, pr: dict, reh: dict) -> bytes:
    doc = _new_doc()
    _title_block(doc, sess, "Post-pitch review: requirements, deck, conversation and client questions")
    _review_sections(doc, reh)
    return _bytes(doc)
