"""Slide visuals: pictures from the pptx, pages from a PDF, chart data as text, and the image blocks sent to the model."""
import io, os, sys, tempfile
from pathlib import Path
os.environ["DATA_DIR"] = tempfile.mkdtemp()
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
import app.main as m
from app import preread

def png(w, h, col):
    b = io.BytesIO(); Image.new("RGB", (w, h), col).save(b, "PNG"); b.seek(0); return b
prs = Presentation()
s1 = prs.slides.add_slide(prs.slide_layouts[5]); s1.shapes.title.text = "Evidence"
s1.shapes.add_picture(png(800, 500, (200, 30, 30)), Inches(1), Inches(2), Inches(5))
s1.shapes.add_picture(png(40, 40, (0, 0, 0)), Inches(7), Inches(2), Inches(0.3))      # a tiny icon: must be dropped
s2 = prs.slides.add_slide(prs.slide_layouts[5]); s2.shapes.title.text = "Results"
cd = CategoryChartData(); cd.categories = ["Q1", "Q2"]; cd.add_series("Sales", (10, 25))
s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), cd)
buf = io.BytesIO(); prs.save(buf); deck = buf.getvalue()
pdf = io.BytesIO(); Image.new("RGB", (1280, 720), (30, 60, 200)).save(pdf, "PDF", save_all=True, append_images=[Image.new("RGB", (1280, 720), (30, 200, 60))])

c = TestClient(m.app)
rfp = "The client needs national coverage and reporting. " * 10
def make(**kw):
    files = {"deck": ("d.pptx", deck, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    if kw.get("pdf"): files["deck_pdf"] = ("d.pdf", pdf.getvalue(), "application/pdf")
    r = c.post("/api/sessions", data={"settings": "{}", "rfp_text": rfp, **({"read_visuals": "1"} if kw.get("vis") else {})}, files=files)
    assert r.status_code == 200, r.text; return r.json()
r = make(); assert r["visuals"]["images"] == 0                                  # tick box off: nothing read
sess = c.get(f"/api/sessions/{r['id']}").json()
assert "[CHART" in sess["slides"][1]["text"] and "25" in sess["slides"][1]["text"], sess["slides"][1]       # chart data becomes text
r = make(vis=True); assert r["visuals"]["source"] == "pptx" and r["visuals"]["images"] == 1, r               # icon dropped
r2 = make(vis=True, pdf=True); assert r2["visuals"]["source"] == "pdf" and r2["visuals"]["images"] == 2 and not r2["visuals"]["warning"], r2
from app import visuals as V
vis = V.load(m._dir(r2["id"]) / "visuals"); assert [v["n"] for v in vis] == [1, 2]
msgs = preread.build_messages("rfp " * 100, [{"n": 1, "title": "a", "text": "", "notes": ""}, {"n": 2, "title": "b", "text": "", "notes": ""}], {}, vis)
kinds = [b["type"] for b in msgs[0]["content"]]; assert kinds.count("image") == 2 and "cache_control" in msgs[0]["content"][-2]
assert [b["type"] for b in preread.build_messages("x", [], {})[0]["content"]] == ["text", "text"]               # no visuals: unchanged
print("VISUALS OK")
